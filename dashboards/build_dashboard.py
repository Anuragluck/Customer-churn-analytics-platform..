"""Build a refreshable Excel dashboard and BI-ready tables from project CSVs."""
from __future__ import annotations

from pathlib import Path

import pandas as pd
from openpyxl import Workbook
from openpyxl.chart import BarChart, DoughnutChart, Reference
from openpyxl.chart.label import DataLabelList
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.worksheet.table import Table, TableStyleInfo


ROOT = Path(__file__).resolve().parents[1]
OUTPUT = ROOT / "dashboards" / "churn_analytics_dashboard.xlsx"
NAVY = "15304B"
TEAL = "12A594"
BLUE = "3B82F6"
LIGHT = "EAF1F7"
WHITE = "FFFFFF"
TEXT = "243447"


def add_table(ws, frame: pd.DataFrame, name: str) -> None:
    ws.append(list(frame.columns))
    for row in frame.itertuples(index=False, name=None):
        ws.append([None if pd.isna(value) else value for value in row])
    if len(frame):
        table = Table(displayName=name, ref=f"A1:{ws.cell(1, len(frame.columns)).column_letter}{len(frame) + 1}")
        table.tableStyleInfo = TableStyleInfo(
            name="TableStyleMedium2", showFirstColumn=False,
            showLastColumn=False, showRowStripes=True, showColumnStripes=False,
        )
        ws.add_table(table)
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions
    for cell in ws[1]:
        cell.font = Font(bold=True, color=WHITE)
        cell.fill = PatternFill("solid", fgColor=NAVY)
        cell.alignment = Alignment(wrap_text=True, vertical="center")
    ws.row_dimensions[1].height = 30
    for column_cells in ws.columns:
        letter = column_cells[0].column_letter
        width = min(max(max(len(str(c.value or "")) for c in column_cells[:250]) + 2, 12), 32)
        ws.column_dimensions[letter].width = width


def main() -> None:
    source = pd.read_csv(ROOT / "data" / "telco_churn.csv")
    scores = pd.read_csv(ROOT / "data" / "active_customer_scores.csv")
    source["Churn"] = source["Churn"].astype(str).str.strip().str.title()
    source["Contract"] = source["Contract"].astype(str).str.strip()
    scores["contract"] = scores["contract"].astype(str).str.strip()

    churn_rows = []
    for contract in ["Month-to-month", "One year", "Two year"]:
        source_group = source[source["Contract"] == contract]
        score_group = scores[scores["contract"] == contract]
        churn_rows.append({
            "Contract": contract,
            "Source customers": int(len(source_group)),
            "Observed churners": int(source_group["Churn"].eq("Yes").sum()),
            "Observed churn rate (%)": round(float(source_group["Churn"].eq("Yes").mean()) * 100, 1),
            "Scored active customers": int(len(score_group)),
            "High-risk customers": int(score_group["is_high_risk"].sum()),
            "Revenue at risk ($/mo)": float(score_group["revenue_at_risk"].sum()),
            "CLV at risk ($)": float(score_group["clv_at_risk"].sum()),
        })
    contracts = pd.DataFrame(churn_rows)

    risk_order = ["Low", "Medium", "High"]
    risk_rows = []
    for level in risk_order:
        group = scores[scores["risk_level"].astype(str).str.title() == level]
        risk_rows.append({
            "Risk level": level,
            "Customers": int(len(group)),
            "Share of scored customers": float(len(group) / max(len(scores), 1)),
            "Average churn probability": float(group["churn_probability"].mean()) if len(group) else 0.0,
            "Revenue at risk ($/mo)": float(group["revenue_at_risk"].sum()),
        })
    risks = pd.DataFrame(risk_rows)

    wb = Workbook()
    dashboard = wb.active
    dashboard.title = "Executive Dashboard"
    dashboard.sheet_view.showGridLines = False
    dashboard.merge_cells("A1:L1")
    dashboard["A1"] = "CUSTOMER CHURN | EXECUTIVE DASHBOARD"
    dashboard["A1"].font = Font(size=20, bold=True, color=WHITE)
    dashboard["A1"].fill = PatternFill("solid", fgColor=NAVY)
    dashboard["A1"].alignment = Alignment(vertical="center")
    dashboard.row_dimensions[1].height = 40
    dashboard.merge_cells("A2:L2")
    dashboard["A2"] = "IBM Telco sample · model scores are illustrative prioritization signals, not realized savings"
    dashboard["A2"].font = Font(italic=True, color=TEXT, size=10)

    churn_rate = float(source["Churn"].eq("Yes").mean())
    metrics = [
        ("SOURCE CUSTOMERS", f"{len(source):,}"),
        ("OBSERVED CHURN RATE", f"{churn_rate:.1%}"),
        ("ACTIVE ROWS SCORED", f"{len(scores):,}"),
        ("EXPECTED MONTHLY REVENUE EXPOSURE", f"${scores['revenue_at_risk'].sum():,.0f}"),
        ("POLICY CONTACT CANDIDATES", f"{int(scores['is_high_risk'].sum()):,}"),
        ("MEAN CALIBRATED RISK", f"{scores['churn_probability'].mean():.1%}"),
    ]
    cards = [("A", "C"), ("E", "G"), ("I", "L"), ("A", "C"), ("E", "G"), ("I", "L")]
    for i, (label, value) in enumerate(metrics):
        top = 4 if i < 3 else 7
        left, right = cards[i]
        dashboard.merge_cells(f"{left}{top}:{right}{top}")
        dashboard.merge_cells(f"{left}{top+1}:{right}{top+1}")
        label_cell = dashboard[f"{left}{top}"]
        value_cell = dashboard[f"{left}{top+1}"]
        label_cell.value = label
        value_cell.value = value
        label_cell.font = Font(bold=True, size=9, color=WHITE)
        label_cell.fill = PatternFill("solid", fgColor=TEAL)
        value_cell.font = Font(bold=True, size=19, color=NAVY)
        value_cell.fill = PatternFill("solid", fgColor=LIGHT)
        label_cell.alignment = value_cell.alignment = Alignment(horizontal="center", vertical="center")
        dashboard.row_dimensions[top].height = 22
        dashboard.row_dimensions[top + 1].height = 35

    contracts_ws = wb.create_sheet("Contract Summary")
    add_table(contracts_ws, contracts, "ContractSummary")
    for row in range(2, contracts_ws.max_row + 1):
        contracts_ws.cell(row, 4).number_format = "0.0"
        for col in (7, 8):
            contracts_ws.cell(row, col).number_format = '$#,##0.00'

    risk_ws = wb.create_sheet("Risk Summary")
    add_table(risk_ws, risks, "RiskSummary")
    for row in range(2, risk_ws.max_row + 1):
        risk_ws.cell(row, 3).number_format = risk_ws.cell(row, 4).number_format = "0.0%"
        risk_ws.cell(row, 5).number_format = '$#,##0.00'

    scores_ws = wb.create_sheet("Customer Scores")
    add_table(scores_ws, scores, "CustomerScores")
    for row in range(2, scores_ws.max_row + 1):
        scores_ws.cell(row, 5).number_format = "0.00%"
        for col in (8, 9, 10):
            scores_ws.cell(row, col).number_format = '$#,##0.00'
    scores_ws.conditional_formatting.add(
        f"E2:E{scores_ws.max_row}",
        __import__("openpyxl").formatting.rule.ColorScaleRule(
            start_type="min", start_color="63BE7B",
            mid_type="percentile", mid_value=50, mid_color="FFEB84",
            end_type="max", end_color="F8696B",
        ),
    )

    source_ws = wb.create_sheet("Source Data")
    add_table(source_ws, source, "SourceCustomers")

    # Doughnut chart for scored risk segmentation.
    risk_chart = DoughnutChart()
    risk_chart.title = "Scored Customers by Risk Level"
    risk_chart.add_data(Reference(risk_ws, min_col=2, min_row=1, max_row=4), titles_from_data=True)
    risk_chart.set_categories(Reference(risk_ws, min_col=1, min_row=2, max_row=4))
    risk_chart.height, risk_chart.width = 7, 11
    risk_chart.dataLabels = DataLabelList()
    risk_chart.dataLabels.showPercent = True
    dashboard.add_chart(risk_chart, "A10")

    churn_chart = BarChart()
    churn_chart.type = "bar"
    churn_chart.style = 10
    churn_chart.title = "Observed Churn Rate by Contract (%)"
    churn_chart.y_axis.title = "Contract"
    churn_chart.x_axis.title = "Observed churn rate (%)"
    churn_chart.x_axis.numFmt = "0.0"
    churn_chart.x_axis.scaling.min = 0
    churn_chart.x_axis.scaling.max = 100
    churn_chart.legend = None
    churn_chart.add_data(Reference(contracts_ws, min_col=4, min_row=1, max_row=4), titles_from_data=True)
    churn_chart.set_categories(Reference(contracts_ws, min_col=1, min_row=2, max_row=4))
    churn_chart.height, churn_chart.width = 7, 12
    churn_chart.dataLabels = DataLabelList()
    churn_chart.dataLabels.showVal = True
    dashboard.add_chart(churn_chart, "G10")

    exposure_chart = BarChart()
    exposure_chart.type = "col"
    exposure_chart.style = 12
    exposure_chart.title = "Expected Monthly Revenue Exposure by Contract"
    exposure_chart.y_axis.title = "Revenue exposure ($/mo)"
    exposure_chart.x_axis.title = "Contract"
    exposure_chart.legend = None
    exposure_chart.add_data(Reference(contracts_ws, min_col=7, min_row=1, max_row=4), titles_from_data=True)
    exposure_chart.set_categories(Reference(contracts_ws, min_col=1, min_row=2, max_row=4))
    exposure_chart.height, exposure_chart.width = 7, 23
    dashboard.add_chart(exposure_chart, "A26")
    for col in range(1, 13):
        dashboard.column_dimensions[__import__("openpyxl").utils.get_column_letter(col)].width = 13
    dashboard.freeze_panes = "A4"

    notes = wb.create_sheet("Definitions")
    notes.append(["Metric", "Definition / use"])
    notes.append(["Observed churn rate", "Share of source sample rows with Churn = Yes; descriptive, not causal."])
    notes.append(["Revenue at risk", "Monthly charges multiplied by calibrated churn probability for scored retained rows."])
    notes.append(["CLV at risk", "Illustrative expected customer value at risk under the project's configured 30-month assumption."])
    notes.append(["High risk", "Flag derived from the cost-selected operating threshold saved with the model."])
    notes.append(["Active rows", "Retained rows in the public sample are used as a teaching proxy; no point-in-time customer snapshots exist."])
    notes.append(["Refresh", "Run `python -m model.score_customers --no-db`, then `python dashboards/build_dashboard.py`."])
    notes.append(["BI tools", "Connect Tableau or Power BI to the PostgreSQL tables, or import the Customer Scores / Source Data tabs."])
    notes.column_dimensions["A"].width = 24
    notes.column_dimensions["B"].width = 110
    notes.freeze_panes = "A2"
    for cell in notes[1]:
        cell.font = Font(bold=True, color=WHITE)
        cell.fill = PatternFill("solid", fgColor=NAVY)
    for row in notes.iter_rows(min_row=2):
        for cell in row:
            cell.alignment = Alignment(wrap_text=True, vertical="top")
    notes.row_dimensions[5].height = 36
    wb.calculation.fullCalcOnLoad = True
    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    wb.save(OUTPUT)
    print(f"Wrote {OUTPUT.relative_to(ROOT)}")
    print(f"Customers: {len(source):,}; active rows scored: {len(scores):,}; monthly exposure: ${scores['revenue_at_risk'].sum():,.2f}")


if __name__ == "__main__":
    main()
